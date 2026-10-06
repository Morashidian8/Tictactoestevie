"""
The live record, read back: which signals fired, in what order, and which won.

`signals_ledger.csv` is the only permanent trace the bot keeps. Everything else
is capped — the in-memory strip holds forty entries, the scorecard holds totals
with no order at all — so a question like "list the rules 1-7 signals of the
last month in order" can only be answered from this file. This reads it.

WHAT THE FILE ALREADY GUARANTEES, and why it can be trusted a month later:

  * a row is written the moment a signal is ISSUED, before anyone knows how it
    ends, so a signal that never settles still leaves a trace;
  * a second row with `win`, `loss` or `void` lands at settlement, and
    `ledger_rows()` lets the later row supersede the earlier one, keyed by the
    window;
  * every append is fsync'd, so a battery death loses nothing already written;
  * a line torn by a power cut is closed off and skipped, costing one record
    rather than the file;
  * windows the bot notices are missing get backfilled.

WHAT IT DOES NOT GUARANTEE: the file lives on the phone and is gitignored, so
it exists in exactly one place. Export it from time to time.

    python ledger_report.py                      # rules 1-7, last 30 days
    python ledger_report.py --days 90 --book all
    python ledger_report.py --book planb
    python ledger_report.py --file exported.csv

Writes ledger_report.xlsx and prints the list. Grading is the bot's own live
grading — what actually happened, not a replay.
"""

import csv
import os
import re
import sys
from collections import defaultdict
from datetime import datetime, timedelta, timezone

OUT = "ledger_report.xlsx"
TEHRAN = timezone(timedelta(hours=3, minutes=30))
LEDGER = os.environ.get("LEDGER_FILE", "signals_ledger.csv")
PLANB = "۸) پلن بی"
FA_DAY = ["دوشنبه", "سه‌شنبه", "چهارشنبه", "پنج‌شنبه", "جمعه", "شنبه", "یکشنبه"]
FA_STATUS = {"win": "برد", "loss": "باخت", "void": "باطل", "issued": "در انتظار"}


def wilson(w, n, z=1.96):
    if not n:
        return 0.0, 0.0
    p = w / n
    d = 1 + z * z / n
    c = p + z * z / (2 * n)
    m = z * ((p * (1 - p) / n + z * z / (4 * n * n)) ** 0.5)
    return max(0.0, (c - m) / d), min(1.0, (c + m) / d)


def read(path):
    """
    Oldest first, one entry per window, later rows superseding earlier ones.

    Deliberately the same collapsing rule the bot itself uses: the settled row
    replaces the issued row it followed. Reading the file without that would
    double-count every signal that reached a conclusion.

    Timestamps outside a believable range are dropped rather than kept. A line
    torn by a power cut does not always fail to parse — `99999999,broken` is a
    perfectly good integer — and one row claiming to be from the year 5000
    would otherwise become the file's "latest", pushing the whole reporting
    window past every real signal and printing an empty month.
    """
    out, torn = {}, 0
    floor = 1_500_000_000                      # 2017; the bot did not exist before
    ceiling = int(datetime.now(timezone.utc).timestamp()) + 86400
    with open(path, newline="") as f:
        for r in csv.DictReader(f):
            try:
                t = int(r["window_epoch"])
            except (TypeError, ValueError, KeyError):
                torn += 1
                continue
            if not (floor <= t <= ceiling):
                torn += 1
                continue
            out[t] = r
    return [out[k] for k in sorted(out)], torn


def split_rules(s):
    """
    The rules named on one row.

    The column is built with `" + ".join(names)`, but two of the rule names
    contain " + " themselves — rule 7 is "باندِ بولینگر + RSI" and rule 2 is
    "۳ حرکتِ هم‌جهت + حرکتِ بزرگ" — so splitting on the separator alone tears
    them in half and invents rules called "RSI" and "حرکتِ بزرگ". Every real
    name begins with a Persian digit and a bracket, or the trophy, so the split
    only fires where one actually starts.
    """
    return [p.strip() for p in re.split(r" \+ (?=[۰-۹]\)|🏆)", s or "")
            if p.strip()]


def is_planb(row):
    """Rule 8 alone — Plan B's own book. Rule 8 never speaks alongside others."""
    return split_rules(row.get("rules")) == [PLANB]


def enrich(rows):
    for r in rows:
        t = int(r["window_epoch"])
        d = datetime.fromtimestamp(t, TEHRAN)
        r["_t"], r["_d"] = t, d
        r["_status"] = (r.get("status") or "").strip()
        r["_won"] = r["_status"] == "win"
        r["_settled"] = r["_status"] in ("win", "loss")
        for k in ("ref", "settle", "delta"):
            try:
                r["_" + k] = float(r.get(k) or 0)
            except ValueError:
                r["_" + k] = 0.0
    return rows


def main():
    argv = sys.argv[1:]
    g = lambda k, d: (argv[argv.index(k) + 1] if k in argv else d)
    days = int(g("--days", "30"))
    bookname = g("--book", "1-7")
    path = g("--file", LEDGER)

    if not os.path.exists(path):
        print(f"{path} not found.")
        print("On the phone it sits next to bot.py. From Telegram, /export "
              "sends you a copy.")
        return

    rows, torn = read(path)
    if not rows:
        print(f"{path} has no readable rows.")
        return
    rows = enrich(rows)

    cut = rows[-1]["_t"] - days * 86400
    rows = [r for r in rows if r["_t"] >= cut]
    keep = {"1-7": lambda r: not is_planb(r),
            "planb": is_planb,
            "all": lambda r: True}.get(bookname)
    if keep is None:
        print("--book must be one of: 1-7, planb, all")
        return
    rows = [r for r in rows if keep(r)]
    if not rows:
        print(f"no signals in the last {days} days for --book {bookname}.")
        return

    label = {"1-7": "قانون‌های ۱ تا ۷", "planb": "پلن بی",
             "all": "همهٔ سیگنال‌ها"}[bookname]
    settled = [r for r in rows if r["_settled"]]
    w = sum(1 for r in settled if r["_won"])
    n = len(settled)
    waiting = sum(1 for r in rows if r["_status"] == "issued")
    void = sum(1 for r in rows if r["_status"] == "void")

    print("=" * 78)
    print(f"  {label} — {days} روزِ گذشته   ({os.path.basename(path)})")
    print("=" * 78)
    print(f"  {rows[0]['_d']:%Y-%m-%d %H:%M} -> {rows[-1]['_d']:%Y-%m-%d %H:%M}"
          f"  (تهران)")
    print(f"  {len(rows):,} سیگنال · {n:,} تسویه‌شده · {waiting:,} در انتظار"
          f" · {void:,} باطل")
    if n:
        lo, hi = wilson(w, n)
        print(f"  {w:,} برد · {n - w:,} باخت · {w / n * 100:.2f}%"
              f"  [{lo * 100:.2f}–{hi * 100:.2f}]")
    if torn:
        print(f"  ⚠️ {torn} خط ناخوانا رد شد (احتمالاً قطع برق هنگام نوشتن)")

    print(f"\n  {'#':>4}  {'تاریخ':<11}{'ساعت':<7}{'روز':<10}"
          f"{'جهت':<7}{'نتیجه':<9}{'حرکت':>10}  قانون‌ها")
    for i, r in enumerate(rows, 1):
        print(f"  {i:>4}  {r['_d']:%Y-%m-%d} {r['_d']:%H:%M}  "
              f"{FA_DAY[r['_d'].weekday()]:<10}"
              f"{(r.get('bet') or ''):<7}"
              f"{FA_STATUS.get(r['_status'], r['_status']):<9}"
              f"{r['_delta']:>+10.2f}  {r.get('rules') or ''}")

    # ---- which rule carried which ------------------------------------------ #
    per = defaultdict(lambda: [0, 0])
    for r in settled:
        for nm in split_rules(r.get("rules")):
            per[nm][0] += 1
            per[nm][1] += 1 if r["_won"] else 0
    if per:
        print(f"\n  {'قانون':<34}{'تعداد':>7}{'برد':>6}{'درصد':>9}")
        for nm in sorted(per, key=lambda k: -per[k][0]):
            tot, won = per[nm]
            print(f"  {nm:<34}{tot:>7,}{won:>6,}{won / tot * 100:>8.1f}%")

    # ---- the workbook ------------------------------------------------------- #
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        print(f"\n(openpyxl not installed — printed only. "
              f"`pip install openpyxl` for {OUT}.)")
        return

    FONT, INK, HEAD_BG = "Arial", "1F2937", "1F3A5F"
    wb = Workbook()
    ws = wb.active
    ws.title = label[:28]
    ws.sheet_view.rightToLeft = True
    head = ["ردیف", "تاریخ (تهران)", "ساعت", "روز", "پنجره (ET)", "جهت",
            "نتیجه", "قیمت آغاز", "قیمت پایان", "حرکت ($)", "قانون‌ها"]
    ws.append(head)
    for c in range(1, len(head) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name=FONT, bold=True, size=10, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEAD_BG)
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
    for i, r in enumerate(rows, 1):
        ws.append([i, f"{r['_d']:%Y-%m-%d}", f"{r['_d']:%H:%M}",
                   FA_DAY[r["_d"].weekday()], r.get("et") or "",
                   r.get("bet") or "",
                   FA_STATUS.get(r["_status"], r["_status"]),
                   r["_ref"], r["_settle"], r["_delta"],
                   r.get("rules") or ""])
    thin = Side(style="thin", color="D1D5DB")
    fills = {"برد": ("DCFCE7", "166534"), "باخت": ("FEE2E2", "991B1B"),
             "در انتظار": ("FEF9C3", "854D0E"), "باطل": ("F3F4F6", "6B7280")}
    for row in ws.iter_rows(min_row=2, max_row=len(rows) + 1, max_col=len(head)):
        for cell in row:
            cell.font = Font(name=FONT, size=10, color=INK)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(bottom=thin)
        for j in (7, 8, 9):
            row[j].number_format = '#,##0.00;(#,##0.00);-'
        row[10].alignment = Alignment(horizontal="right", vertical="center")
        bg, fg = fills.get(row[6].value, ("FFFFFF", INK))
        row[6].fill = PatternFill("solid", fgColor=bg)
        row[6].font = Font(name=FONT, size=10, bold=True, color=fg)
    for i, wdt in enumerate((7, 14, 8, 11, 20, 9, 11, 13, 13, 12, 46), 1):
        ws.column_dimensions[get_column_letter(i)].width = wdt
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(head))}{len(rows) + 1}"

    sm = wb.create_sheet("خلاصه")
    sm.sheet_view.rightToLeft = True
    sm["A1"] = f"{label} — {days} روزِ گذشته"
    sm["A1"].font = Font(name=FONT, bold=True, size=14, color=INK)
    body = [("", ""), ("بازه", f"{rows[0]['_d']:%Y-%m-%d} تا "
                               f"{rows[-1]['_d']:%Y-%m-%d}"),
            ("کل سیگنال‌ها", f"{len(rows):,}"),
            ("تسویه‌شده", f"{n:,}"), ("در انتظار", f"{waiting:,}"),
            ("باطل", f"{void:,}")]
    if n:
        lo, hi = wilson(w, n)
        body += [("برد", f"{w:,}"), ("باخت", f"{n - w:,}"),
                 ("درصد برد", f"{w / n * 100:.2f}%"),
                 ("بازهٔ اطمینان ۹۵٪",
                  f"{lo * 100:.2f}% تا {hi * 100:.2f}%")]
    body += [("", ""), ("منبع", "رکوردِ زندهٔ خودِ بات — بازپخش نیست."),
             ("فایل", os.path.basename(path))]
    r_ = 3
    for a, b in body:
        sm.cell(row=r_, column=1, value=a).font = Font(name=FONT, size=10,
                                                       bold=bool(a), color=INK)
        c = sm.cell(row=r_, column=2, value=b)
        c.font = Font(name=FONT, size=10, color=INK)
        c.alignment = Alignment(horizontal="right", wrap_text=True)
        r_ += 1
    sm.column_dimensions["A"].width = 24
    sm.column_dimensions["B"].width = 54

    if per:
        rs = wb.create_sheet("به تفکیک قانون")
        rs.sheet_view.rightToLeft = True
        rs.append(["قانون", "تعداد", "برد", "باخت", "درصد"])
        for c in range(1, 6):
            cell = rs.cell(row=1, column=c)
            cell.font = Font(name=FONT, bold=True, size=10, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=HEAD_BG)
            cell.alignment = Alignment(horizontal="center")
        for nm in sorted(per, key=lambda k: -per[k][0]):
            tot, won = per[nm]
            rs.append([nm, tot, won, tot - won, won / tot])
        for row in rs.iter_rows(min_row=2, max_row=len(per) + 1, max_col=5):
            for cell in row:
                cell.font = Font(name=FONT, size=10, color=INK)
                cell.alignment = Alignment(horizontal="center")
                cell.border = Border(bottom=thin)
            row[4].number_format = "0.0%"
        for i, wdt in enumerate((40, 9, 8, 8, 10), 1):
            rs.column_dimensions[get_column_letter(i)].width = wdt
        rs.freeze_panes = "A2"

    wb.save(OUT)
    print(f"\nwrote {OUT} — {len(rows):,} rows")


if __name__ == "__main__":
    main()
